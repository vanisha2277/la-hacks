#!/usr/bin/env python3
"""Build the Unit 200 valve list in the same column layout as the example KKS valve list.

  python3 build_valve_list.py      writes valve_list/unit200_valve_list.xlsx and .json

Values we have no record of are 'N/A' or 'LATER', the same way the example sheet does it.
That is on purpose: an empty cell is a gap the agent should report, not fill in.
The valve list reflects the *official* record at Rev 5, so it still says PSV set pressure 17.2 barg
and LV fail-closed (see contradictions C01 and C02).
"""
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import kks

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "valve_list")

COLUMNS = ["No", "Rev", "Status", "REVISION NOTE", "P&ID", "SYSTEM", "TYPE", "KKS TAG", "UNIT NO", "SYSTEM CODE",
           "SYSTEM SEQ NO", "ITEM CODE", "FIRST SEQ NO", "Description", "Fluid",
           "Normal Mass Flow (kg/s)", "Normal Pressure Inlet (Barg)", "Normal Temp (Deg.C)",
           "Design Mass Flow (kg/s)", "Design Pressure Inlet (Barg)", "Design Temp (Deg.C)",
           "Set Pressure (BARG)", "Back Pressure (BARG)", "PDC", "Line Size", "Line Material", "Line Spec",
           "ACTUATOR TYPE", "VALVE TYPE", "Fail Position", "Travel Time", "Leakage Class", "SIZE(DN)",
           "BODY MATERIAL", "TRIM MATERIAL", "SEAT MATERIAL", "VALVE RATING", "END TYPE", "Action of Actuator",
           "VALVE MAKER", "VALVE MODEL NO", "ACTUATOR MAKER", "ACTUATOR MODEL NO", "REMARK", "ACCESSORIES"]

NA = "N/A"
LATER = "LATER"

# per valve: P&ID, last drawing rev, status, revision note, fluid, DN, actuator, valve type, fail, action, set pressure, remark, accessories
V = {
    "FV-101": ("PID-100", "A", "MODIFY", "", "C3/C4 MIXED FEED", 100, "PNEUMATIC ACTUATOR", "GLOBE", "Fail Close", "SINGLE ACTION", NA, "", "POSITIONER"),
    "BV-101": ("PID-100", "3", "NEW", "ADDED BY MOC-2022-007", "C3/C4 MIXED FEED", 100, "MANUAL", "BALL", "-", "-", NA, "E-101 BYPASS", "LOCKING DEVICE"),
    "FV-201": ("PID-200", "2", "MODIFY", "TRIM UPSIZED BY MOC-2020-014", "PROPANE (C3)", 80, "PNEUMATIC ACTUATOR", "GLOBE", "Fail Close", "SINGLE ACTION", NA, "CV 45 TO CV 75 AT REV 2", "POSITIONER"),
    "LV-202": ("PID-200", "5", "MODIFY", "FAIL POSITION FO TO FC BY MOC-2024-031", "PROPANE (C3)", 50, "PNEUMATIC ACTUATOR", "GLOBE", "Fail Close", "SINGLE ACTION", NA, "", "POSITIONER"),
    "PV-201": ("PID-200", "A", "MODIFY", "", "COOLING WATER", 150, "PNEUMATIC ACTUATOR", "BUTTERFLY", "Fail Open", "SINGLE ACTION", NA, "", "POSITIONER"),
    "TV-203": ("PID-200", "A", "MODIFY", "", "STEAM", 80, "PNEUMATIC ACTUATOR", "GLOBE", "Fail Close", "SINGLE ACTION", NA, "", "POSITIONER"),
    "PSV-201": ("PID-200", "A", "MODIFY", "", "C3/C4 VAPOR", 150, "SELF ACTUATING", "SAFETY & RELIEF", "-", "SELF ACTUATOR", 17.2, "SET PRESSURE PER DRAWING REV 5 AND DATASHEET", ""),
    "CV-102A": ("PID-200", "1", "NEW", "ADDED BY MOC-2019-021", "PROPANE (C3)", 80, "SELF ACTUATING", "SWING CHECK", "-", "SELF ACTUATOR", NA, "REFLUX PUMP A DISCHARGE", ""),
    "CV-102B": ("PID-200", "1", "NEW", "ADDED BY MOC-2019-021", "PROPANE (C3)", 80, "SELF ACTUATING", "SWING CHECK", "-", "SELF ACTUATOR", NA, "REFLUX PUMP B DISCHARGE", ""),
}

PID_DOC = {"PID-100": "PID-100 SHEET 1/2", "PID-200": "PID-200 SHEET 2/2"}


def rows():
    reg = {r["legacy"]: r for r in kks.REGISTER}
    out = []
    for i, (legacy, v) in enumerate(V.items(), start=1):
        r = reg[legacy]
        pid, rev, status, note, fluid, dn, act, vtype, fail, action, setp, remark, acc = v
        out.append({
            "No": i, "Rev": rev, "Status": status, "REVISION NOTE": note, "P&ID": PID_DOC[pid], "SYSTEM": r["system"],
            "TYPE": r["type"], "KKS TAG": r["kks"], "UNIT NO": r["unit_no"], "SYSTEM CODE": r["system_code"],
            "SYSTEM SEQ NO": r["system_seq"], "ITEM CODE": r["item_code"], "FIRST SEQ NO": r["first_seq"],
            "Description": r["description"], "Fluid": fluid,
            "Normal Mass Flow (kg/s)": NA, "Normal Pressure Inlet (Barg)": NA, "Normal Temp (Deg.C)": NA,
            "Design Mass Flow (kg/s)": NA, "Design Pressure Inlet (Barg)": NA, "Design Temp (Deg.C)": NA,
            "Set Pressure (BARG)": setp, "Back Pressure (BARG)": NA, "PDC": "LFC1" if r["system_code"] == "LFC" else "LFB1",
            "Line Size": dn, "Line Material": "A106-B", "Line Spec": "1C1", "ACTUATOR TYPE": act, "VALVE TYPE": vtype,
            "Fail Position": fail, "Travel Time": LATER if act.startswith("PNEU") else "-", "Leakage Class": LATER,
            "SIZE(DN)": dn, "BODY MATERIAL": "A216-WCB", "TRIM MATERIAL": LATER, "SEAT MATERIAL": LATER,
            "VALVE RATING": "300#", "END TYPE": "Flanged - RF", "Action of Actuator": action,
            "VALVE MAKER": LATER, "VALVE MODEL NO": LATER, "ACTUATOR MAKER": LATER, "ACTUATOR MODEL NO": LATER,
            "REMARK": remark, "ACCESSORIES": acc,
        })
    return out


def main():
    os.makedirs(OUT, exist_ok=True)
    data = rows()
    json.dump(data, open(os.path.join(OUT, "unit200_valve_list.json"), "w"), indent=2)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ball, Gate, Globe, ETC"
    ws["A1"] = "UNIT 200 VALVE LIST  -  SYNTHETIC DEMO DATA, NOT A REAL PLANT  (drawing revision 5)"
    ws["A1"].font = Font(bold=True)
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=c, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(data, start=3):
        for c, name in enumerate(COLUMNS, start=1):
            ws.cell(row=r, column=c, value=row[name])
    for c, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(10, min(34, len(name) + 2))
    ws.freeze_panes = "I3"
    wb.save(os.path.join(OUT, "unit200_valve_list.xlsx"))
    print("wrote %d valves to %s" % (len(data), OUT))


if __name__ == "__main__":
    main()
